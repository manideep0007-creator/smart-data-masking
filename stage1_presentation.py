from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = r"D:\Innovex\smart-data-masking\Smart_Data_Masking_Stage1.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(10, 16, 31)
PANEL = RGBColor(20, 30, 51)
CYAN = RGBColor(0, 220, 240)
WHITE = RGBColor(242, 247, 255)
MUTED = RGBColor(169, 187, 211)
GREEN = RGBColor(38, 214, 151)
PURPLE = RGBColor(173, 111, 255)


def bg(slide):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = NAVY
    # Accent rail
    rail = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.12), prs.slide_height)
    rail.fill.solid(); rail.fill.fore_color.rgb = CYAN; rail.line.fill.background()


def text(slide, x, y, w, h, value, size=18, color=WHITE, bold=False, font="Aptos", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None: p.alignment = align
    run = p.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return box


def panel(slide, x, y, w, h, fill=PANEL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(41, 59, 88)
    return shape


def header(slide, label, title, subtitle=None):
    text(slide, .55, .35, 12, .28, label.upper(), 11, CYAN, True)
    text(slide, .55, .72, 12, .55, title, 29, WHITE, True)
    if subtitle: text(slide, .58, 1.35, 12, .45, subtitle, 14, MUTED)


def bullet(slide, x, y, w, value, color=WHITE, size=17):
    text(slide, x, y, .25, .32, "•", size, CYAN, True)
    text(slide, x+.32, y, w-.32, .6, value, size, color)

# Slide 1
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
text(s, .8, .65, 11.8, .35, "HACKATHON · STAGE 1", 13, CYAN, True)
text(s, .8, 1.45, 11.6, 1.35, "Smart Data\nMasking System", 38, WHITE, True)
text(s, .85, 4.05, 10.5, .65, "Protect sensitive fields while teams use customer data", 22, MUTED)
panel(s, .85, 5.35, 5.5, .82)
text(s, 1.12, 5.58, 5.0, .3, "PII SCAN  →  ROLE POLICY  →  MASKED RESULTS", 14, CYAN, True)
text(s, .85, 6.8, 11.5, .3, "A working prototype by Manideep & teammate", 12, MUTED)

# Slide 2
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
header(s, "01 · The challenge", "Sensitive data needs different levels of access", "Support and analytics teams need useful data—but not always the raw personal details.")
panel(s, .75, 2.15, 5.75, 3.6)
text(s, 1.08, 2.48, 5.0, .4, "THE RISK", 13, CYAN, True)
bullet(s, 1.08, 3.15, 5.0, "Too much internal access can expose personal and financial fields.")
bullet(s, 1.08, 4.15, 5.0, "Manual masking is hard to apply consistently across queries.")
panel(s, 6.8, 2.15, 5.75, 3.6)
text(s, 7.15, 2.48, 5.0, .4, "OUR GOAL", 13, GREEN, True)
bullet(s, 7.15, 3.15, 4.9, "Detect likely sensitive columns in a sample database.")
bullet(s, 7.15, 4.15, 4.9, "Return different views by role without changing stored records.")
text(s, .8, 6.55, 11.7, .4, "Prototype scope: a local SQLite demo, accessed through our API gateway.", 14, MUTED)

# Slide 3
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
header(s, "02 · The solution", "One query. Three role-based views.", "Masking is applied to returned query results in memory.")
roles = [("ADMIN", "Full values", "4111-2222-3333-1111", RGBColor(255, 91, 102)), ("SUPPORT", "Partial masking", "4111-XXXX-XXXX-1111", CYAN), ("ANALYST", "Redacted", "[REDACTED]", PURPLE)]
for i,(name,desc,result,c) in enumerate(roles):
    x=.78+i*4.2
    panel(s,x,2.3,3.85,2.35)
    text(s,x+.25,2.62,3.3,.32,name,14,c,True)
    text(s,x+.25,3.12,3.3,.3,desc,16,WHITE,True)
    text(s,x+.25,3.72,3.35,.5,result,15,MUTED)
text(s, .9, 5.35, 11.5, .55, "Same SELECT query  →  backend checks role  →  result is masked before it reaches the dashboard", 17, WHITE, True)
text(s, .9, 6.25, 11.5, .35, "Sample records are synthetic. The database remains unchanged by masking.", 13, MUTED)

# Slide 4
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
header(s, "03 · Architecture", "A lightweight query gateway", "The dashboard never applies masking itself; it asks the backend for protected results.")
steps=[("01", "Dashboard", "Role + SQL query"), ("02", "FastAPI", "/health · /scan · /query"), ("03", "SQLite", "Synthetic customers"), ("04", "Detect + mask", "Heuristics + role rules"), ("05", "Results", "Columns + masked rows")]
for i,(num,title,desc) in enumerate(steps):
    x=.65+i*2.55
    panel(s,x,2.7,2.22,1.7)
    text(s,x+.18,2.92,.5,.3,num,12,CYAN,True)
    text(s,x+.18,3.35,1.9,.35,title,16,WHITE,True)
    text(s,x+.18,3.82,1.95,.42,desc,11,MUTED)
    if i<4: text(s,x+2.25,3.3,.3,.4,"→",22,CYAN,True)
panel(s, .85, 5.25, 11.55, .82, RGBColor(15, 28, 43))
text(s, 1.1, 5.49, 11.0, .32, "TECH: Python · FastAPI · SQLite · HTML/CSS/JavaScript UI · pytest", 15, WHITE, True)
text(s, .9, 6.55, 11.5, .35, "Use the verified UI/commands for your final demo; this slide describes the architecture at a high level.", 11, MUTED)

# Slide 5
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
header(s, "04 · Smart detection", "What is—and isn’t—AI in our prototype?", "We want to be precise about the current implementation.")
panel(s, .8, 2.15, 5.65, 3.55)
text(s, 1.12, 2.48, 4.9, .35, "CURRENTLY IMPLEMENTED", 13, CYAN, True)
bullet(s, 1.12, 3.15, 4.95, "Rules look at column names and sample-value patterns.")
bullet(s, 1.12, 4.1, 4.95, "Rules assign likely PII/financial tags, then role policy masks results.")
panel(s, 6.85, 2.15, 5.65, 3.55)
text(s, 7.18, 2.48, 4.9, .35, "HONEST AI STATEMENT", 13, PURPLE, True)
text(s, 7.18, 3.15, 4.85, 1.4, "The current detector is heuristic automation—not a trained ML model or generative AI feature.", 18, WHITE, True)
text(s, 7.18, 4.85, 4.85, .5, "AI tools may have helped us build the prototype; they are not the product’s detector.", 13, MUTED)

# Slide 6
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
header(s, "05 · Progress & next steps", "Working MVP, with clear boundaries", "We built and tested the core demonstration flow.")
items=[("WORKING", "Scan sample columns; role-based query masking; read-only query checks; dashboard flow."), ("LIMITS", "Heuristic detection can miss data; demo roles are not authenticated; local SQLite only."), ("NEXT", "Verified identity and authorization; broader database support; stronger detection and audit controls.")]
for i,(title,desc) in enumerate(items):
    y=2.15+i*1.35
    panel(s,.85,y,11.55,1.05)
    c=[GREEN,CYAN,PURPLE][i]
    text(s,1.12,y+.18,1.7,.3,title,13,c,True)
    text(s,2.8,y+.16,9.15,.65,desc,15,WHITE)
text(s, .95, 6.55, 11.3, .38, "This is a hackathon prototype; it does not guarantee GDPR, HIPAA, or PCI-DSS compliance.", 13, MUTED, True)

# Slide 7
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
text(s, .85, .65, 11.5, .35, "LIVE DEMO · 3 MINUTES", 13, CYAN, True)
text(s, .85, 1.15, 11.6, .65, "Let’s see the same data through different roles.", 29, WHITE, True)
for i,(n,head,detail) in enumerate([
    ("1", "Scan", "Point out card_number, SSN, email and sensitivity tags."),
    ("2", "Compare", "Run one query as Admin → Support → Analyst."),
    ("3", "Protect", "If available, show a write query being rejected."),
]):
    y=2.35+i*1.15
    panel(s, .9, y, 11.4, .85)
    text(s,1.17,y+.2,.45,.35,n,16,CYAN,True)
    text(s,1.8,y+.18,1.7,.35,head,16,WHITE,True)
    text(s,3.45,y+.18,8.35,.45,detail,14,MUTED)
text(s, .95, 6.25, 11.4, .6, "Thank you · Questions?", 23, GREEN, True)

# Footer slide numbers
for i, slide in enumerate(prs.slides, 1):
    text(slide, 12.15, 7.1, .55, .2, f"{i:02d} / {len(prs.slides):02d}", 9, MUTED, False, align=PP_ALIGN.RIGHT)

prs.save(OUT)
print(OUT)
