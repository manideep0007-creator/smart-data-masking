import os
from pptx import Presentation
from pptx.util import Pt

src_path = r"C:\Users\N.Manideep\.gemini\antigravity\brain\9f9d30c4-c18e-45b1-ada1-18cc3f8ba5fb\.user_uploaded\media_1793360423704_4632940e.pptx"
prs = Presentation(src_path)

def set_para_text(p, new_text):
    if not p.runs:
        p.text = new_text
        return
    # Preserve font properties of the first run
    font_name = p.runs[0].font.name
    font_size = p.runs[0].font.size
    font_bold = p.runs[0].font.bold
    try:
        font_color = p.runs[0].font.color.rgb
    except Exception:
        font_color = None

    p.text = new_text
    if p.runs:
        if font_name: p.runs[0].font.name = font_name
        if font_size: p.runs[0].font.size = font_size
        if font_bold is not None: p.runs[0].font.bold = font_bold
        if font_color: p.runs[0].font.color.rgb = font_color

# ====================================================
# SLIDE 3: Solution & Gateway
# ====================================================
s3 = prs.slides[2]
# Shape 4: Connect description
set_para_text(s3.shapes[4].text_frame.paragraphs[0], 
              "Database Abstraction Layer. SQLite native authorizer & PostgreSQL read-only session controls.")

# Shape 10: Control description
set_para_text(s3.shapes[10].text_frame.paragraphs[0], 
              "JWT-Verified RBAC policies. Cryptographic token verification for Admin, Support, Analyst.")

# Shape 15: Table - Analyst row
table = s3.shapes[15].table
table.cell(3, 0).text = "Analyst\nFull Privacy / Hash"
table.cell(3, 2).text = "sha256:37b95487"
table.cell(3, 3).text = "sha256:6b2079b1"

# ====================================================
# SLIDE 4: Architecture
# ====================================================
s4 = prs.slides[3]
# Shape 5: Database box
set_para_text(s4.shapes[5].text_frame.paragraphs[0], "Database Abstraction")
set_para_text(s4.shapes[6].text_frame.paragraphs[0], "SQLite & PostgreSQL (100% unmodified)")

# Shape 14: Driver authorizer
set_para_text(s4.shapes[14].text_frame.paragraphs[0], "Driver-Level Read-Only Guard")
set_para_text(s4.shapes[15].text_frame.paragraphs[0], "SQLite set_authorizer & PostgreSQL session read-only.")

# Shape 22: Masking engine
set_para_text(s4.shapes[22].text_frame.paragraphs[0], "In-Memory Masking & Hashing Engine")
set_para_text(s4.shapes[23].text_frame.paragraphs[0], "Targeted partial masking & deterministic HMAC-SHA256.")

# Endpoints boxes
set_para_text(s4.shapes[25].text_frame.paragraphs[0], "JWT Auth & /health")
set_para_text(s4.shapes[26].text_frame.paragraphs[0], "Bcrypt password verification & signed JWT tokens.")

set_para_text(s4.shapes[27].text_frame.paragraphs[0], "GET /scan & /scan/start")
set_para_text(s4.shapes[28].text_frame.paragraphs[0], "Non-blocking background scanner & metadata registry.")

# ====================================================
# SLIDE 5: Security Engineering
# ====================================================
s5 = prs.slides[4]
set_para_text(s5.shapes[6].text_frame.paragraphs[0], 
              "SQL validator blocks comments (-- and /* */), rejects write/DDL keywords, and permits safe CTEs.")
set_para_text(s5.shapes[10].text_frame.paragraphs[0], "Database Kernel Protection")
set_para_text(s5.shapes[11].text_frame.paragraphs[0], 
              "SQLite authorizer & PostgreSQL read-only transaction safeguards enforce zero writes.")

# ====================================================
# SLIDE 6: Why PII Shield Wins & Stats
# ====================================================
s6 = prs.slides[5]
# 15/15 -> 39/39
set_para_text(s6.shapes[9].text_frame.paragraphs[0], "39/39")
set_para_text(s6.shapes[10].text_frame.paragraphs[0], "Automated unit tests passed (pytest)")

# 2 -> 34/34 E2E
set_para_text(s6.shapes[15].text_frame.paragraphs[0], "34/34")
set_para_text(s6.shapes[16].text_frame.paragraphs[0], "Live end-to-end integration tests passed")

# ====================================================
# SLIDE 7: Live Prototype Demo & Next Steps
# ====================================================
s7 = prs.slides[6]
# Clean up quote encoding on Shape 16
set_para_text(s7.shapes[16].text_frame.paragraphs[0], '"DROP TABLE customers" blocked with a 400 Bad Request.')

# Update Role Switching to JWT Persona Switcher
set_para_text(s7.shapes[11].text_frame.paragraphs[0], "JWT Persona Switcher & Hashing")
set_para_text(s7.shapes[12].text_frame.paragraphs[0], "Live cryptographic JWT personas with Redact vs. Deterministic Hash.")

# Update Next Steps (Since JWT and PostgreSQL are ALREADY built!)
set_para_text(s7.shapes[18].text_frame.paragraphs[0], "Next steps · Stage 2 (Future Roadmap)")
set_para_text(s7.shapes[19].text_frame.paragraphs[0], "ML/NLP (Microsoft Presidio) for unstructured document PII")
set_para_text(s7.shapes[19].text_frame.paragraphs[1], "Real-time Kafka streaming query & event masking")
set_para_text(s7.shapes[19].text_frame.paragraphs[2], "Enterprise SSO & automated regulatory compliance reports")

# Save outputs
target1 = r"D:\Innovex\smart-data-masking\Smart_Data_Masking_PII_Shield.pptx"
target2 = r"D:\Innovex\smart-data-masking\Smart_Data_Masking_Stage1.pptx"

prs.save(target1)
prs.save(target2)
print("Updated presentations saved successfully:")
print("1.", target1)
print("2.", target2)

