from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = r"D:\Innovex\smart-data-masking\PII_Shield_INNOVEX_BFSI_Hackathon.pptx"
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
NAVY=RGBColor(8,12,24); PANEL=RGBColor(15,23,42); PANEL2=RGBColor(20,31,54)
CYAN=RGBColor(0,240,255); GREEN=RGBColor(16,185,129); AMBER=RGBColor(245,158,11)
WHITE=RGBColor(255,255,255); MUTED=RGBColor(163,180,205); RED=RGBColor(255,99,112)

def shape(slide,x,y,w,h,color,kind=MSO_SHAPE.ROUNDED_RECTANGLE,line=None):
    s=slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=color
    if line: s.line.color.rgb=line
    else: s.line.fill.background()
    return s

def txt(slide,x,y,w,h,value,size=16,color=WHITE,bold=False,align=None,font="Aptos"):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(.04); tf.margin_right=Inches(.04); tf.margin_top=Inches(.02); tf.margin_bottom=Inches(.02)
    p=tf.paragraphs[0]
    if align is not None: p.alignment=align
    r=p.add_run(); r.text=value; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def base(kicker,title,sub=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
    shape(s,0,0,.12,7.5,CYAN,MSO_SHAPE.RECTANGLE)
    txt(s,.62,.34,12,.25,kicker.upper(),10,CYAN,True)
    txt(s,.62,.72,12,.62,title,28,WHITE,True)
    if sub: txt(s,.66,1.4,11.9,.42,sub,13,MUTED)
    return s

def bullet(s,x,y,w,heading,detail,accent=CYAN):
    shape(s,x,y,.055,.68,accent,MSO_SHAPE.RECTANGLE)
    txt(s,x+.18,y,w-.2,.28,heading,15,WHITE,True)
    txt(s,x+.18,y+.31,w-.2,.4,detail,11,MUTED)

def footer(s,n):
    txt(s,.65,7.13,11,.18,"PII SHIELD  •  AI FRONTIER CHALLENGE 2026  •  INNOSMART",8,MUTED)
    txt(s,12.1,7.1,.55,.2,f"{n:02d}/07",8,MUTED,False,PP_ALIGN.RIGHT)

# 1 Cover
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
shape(s,0,0,.14,7.5,CYAN,MSO_SHAPE.RECTANGLE)
# decorative security motif
for i,c in enumerate([CYAN,GREEN,AMBER]): shape(s,9.45+i*.48,1.0,.24,.24,c,MSO_SHAPE.OVAL)
txt(s,.82,.65,8,.3,"AI FRONTIER CHALLENGE 2026  |  INNOSMART  |  BFSI TRACK · PROBLEM 11",11,CYAN,True)
txt(s,.82,1.55,8.5,1.35,"PII SHIELD",42,WHITE,True)
txt(s,.85,3.02,8.4,.9,"Smart Data Masking System\nfor PII Protection",21,MUTED,True)
shape(s,.86,4.25,.07,.9,GREEN,MSO_SHAPE.RECTANGLE)
txt(s,1.08,4.28,8.2,.85,"Autonomous In-Flight PII Discovery &\nDynamic Role-Based Query Masking",17,WHITE,True)
shape(s,9.45,2.05,2.65,2.65,PANEL,MSO_SHAPE.HEXAGON,CYAN)
txt(s,9.8,2.85,1.95,.7,"PII\nSHIELD",20,CYAN,True,PP_ALIGN.CENTER)
txt(s,.88,5.78,11.6,.3,"TEAM INNOVEX",14,GREEN,True)
txt(s,.88,6.14,11.6,.55,"Team Leader: Mantri Praneeth   •   Mangu Venkata Sai Pravallika\nMajji Chaitanya Sai   •   Nallabanda Manideep",11,WHITE)
footer(s,1)

# 2 Problem
s=base("01 · Problem statement","Live database queries can expose sensitive financial PII","Teams need data to serve customers and analyze operations—but broad raw-data access increases exposure.")
items=[("EXCESSIVE PRIVILEGE","Internal users may receive more raw customer data than their task requires.",RED),("STATIC PROTECTION","Static masking or separate extracts may not adapt to the user and query context.",AMBER),("PRIVACY & COMPLIANCE RISK","Uncontrolled access to identifiers and financial fields can increase privacy and governance risk.",CYAN)]
for i,(h,d,c) in enumerate(items):
    x=.8+i*4.15; shape(s,x,2.35,3.75,2.5,PANEL,line=RGBColor(38,55,82)); shape(s,x+.28,2.7,.5,.08,c,MSO_SHAPE.RECTANGLE)
    txt(s,x+.28,3.02,3.2,.5,h,14,c,True); txt(s,x+.28,3.7,3.15,.85,d,15,WHITE)
shape(s,.82,5.35,11.7,.85,PANEL2)
txt(s,1.1,5.6,11,.3,"DESIGN GOAL   Apply least-privilege views dynamically, while leaving source records unchanged.",15,GREEN,True)
footer(s,2)

# 3 solution table
s=base("02 · Proposed solution","One query. Different views by role.","The gateway applies a role policy to returned values; the demo uses synthetic customer data.")
cols=[.8,3.0,5.65,8.55,10.75]; widths=[2.2,2.65,2.9,2.2,1.8]
headers=["ROLE","CUSTOMER NAME","CREDIT CARD","SSN","SALARY"]
y=2.12
for x,w,h in zip(cols,widths,headers): shape(s,x,y,w,.52,RGBColor(27,42,67)); txt(s,x+.1,y+.14,w-.2,.22,h,10,CYAN,True)
rows=[("Admin · Full access","Asha","4111-2222-3333-1111","123-45-6789","$85,000",RED),("Support · Targeted","Asha","4111-XXXX-XXXX-1111","XXX-XX-6789","[CONFIDENTIAL]",GREEN),("Analyst · Privacy","[REDACTED]","[REDACTED]","[REDACTED]","[REDACTED]",AMBER)]
for ri,row in enumerate(rows):
    y=2.77+ri*.88; fill=PANEL if ri%2==0 else PANEL2
    for ci,(x,w,val) in enumerate(zip(cols,widths,row[:5])):
        shape(s,x,y,w,.7,fill,line=RGBColor(36,51,74)); txt(s,x+.1,y+.2,w-.2,.3,val,10,row[5] if ci==0 else WHITE,ci==0)
shape(s,.82,5.75,11.6,.65,RGBColor(12,35,43))
txt(s,1.05,5.94,11,.25,"Same SQL request  →  role policy  →  masked response; database values are not rewritten.",13,GREEN,True)
footer(s,3)

# 4 architecture
s=base("03 · Technical architecture","A lightweight application-layer query gateway","Tech stack: Python · FastAPI · SQLite · Vanilla HTML/CSS/JavaScript Security Operations Console")
steps=[("01","WEB CONSOLE","Role + SQL"),("02","FASTAPI GATEWAY","Validate · scan · mask"),("03","SQLITE DB","Synthetic records"),("04","POLICY ENGINE","Classify · transform"),("05","SAFE RESPONSE","Columns + rows")]
for i,(n,h,d) in enumerate(steps):
    x=.55+i*2.58; shape(s,x,2.55,2.12,1.62,PANEL,line=RGBColor(38,55,82)); txt(s,x+.18,2.75,.4,.25,n,10,CYAN,True); txt(s,x+.18,3.12,1.8,.32,h,12,WHITE,True); txt(s,x+.18,3.58,1.8,.4,d,10,MUTED)
    if i<4: txt(s,x+2.18,3.1,.35,.4,"›",24,CYAN,True,PP_ALIGN.CENTER)
for i,(endpoint,desc) in enumerate([("/health","Service + DB status"),("/scan","Heuristic schema scan"),("/query","Role-masked SELECT")]):
    x=.9+i*4.05; shape(s,x,4.85,3.65,.82,PANEL2); txt(s,x+.2,5.03,1.2,.28,endpoint,14,CYAN,True); txt(s,x+1.25,5.06,2.2,.3,desc,11,WHITE)
text=s if False else None
footer(s,4)

# 5 Security
s=base("04 · Security engineering","Dual-layer guard for read-only query execution","In-flight masking changes the response—not the source database records.")
for i,(title,desc,c) in enumerate([
    ("APPLICATION LAYER","SQL validator permits read queries and rejects write/administrative statements and chained statements.",CYAN),
    ("SQLITE ENGINE LAYER","SQLite authorizer denies non-read operations at the database connection level.",GREEN)]):
    x=.85+i*6.05; shape(s,x,2.28,5.55,2.2,PANEL,line=RGBColor(38,55,82)); txt(s,x+.32,2.62,4.9,.35,title,14,c,True); txt(s,x+.32,3.22,4.8,.85,desc,15,WHITE)
shape(s,.9,5.08,3.55,.82,PANEL2); txt(s,1.12,5.34,3.1,.27,"ZERO SOURCE WRITES",13,GREEN,True)
shape(s,4.88,5.08,3.55,.82,PANEL2); txt(s,5.12,5.34,3.1,.27,"SYNTHETIC DEMO DATA",13,CYAN,True)
shape(s,8.86,5.08,3.55,.82,PANEL2); txt(s,9.08,5.34,3.1,.27,"API IS THE GATEWAY",13,AMBER,True)
txt(s,.95,6.28,11.3,.4,"Limit: direct database-file access bypasses application-layer masking; this prototype is not a compliance certification.",11,MUTED)
footer(s,5)

# 6 scorecard
s=base("05 · Validation","Core behaviors verified for the demo","Reported validation results from the project test run and live role walkthrough.")
stats=[("15 / 15","UNIT TESTS",CYAN),("3","ROLE POLICIES",GREEN),("2","QUERY GUARD LAYERS",AMBER),("0","DATABASE MUTATIONS",CYAN)]
for i,(big,label,c) in enumerate(stats):
    x=.8+i*3.12; shape(s,x,2.55,2.72,1.8,PANEL,line=RGBColor(38,55,82)); txt(s,x+.2,2.92,2.3,.65,big,30,c,True,PP_ALIGN.CENTER); txt(s,x+.2,3.72,2.3,.3,label,10,WHITE,True,PP_ALIGN.CENTER)
shape(s,.9,5.1,11.45,.95,PANEL2)
txt(s,1.18,5.4,10.9,.35,"Validated demo flow: health check · PII scan · Admin/Support/Analyst comparison · rejected write query",13,WHITE,True)
txt(s,.95,6.35,11.3,.32,"Use these numbers only if your team personally confirms the test run and live behavior before presenting.",10,AMBER,True)
footer(s,6)

# 7 demo
s=base("06 · Live prototype","PII Shield in action","Three-minute walkthrough · use the actual running Security Operations Console")
demos=[("01","HEALTH CHECK","Show gateway and database connection status."),("02","PII SCHEMA SCAN","Identify card_number, SSN, email and sensitivity labels."),("03","ROLE COMPARISON","Run the same SELECT as Admin, Support, then Analyst."),("04","WRITE REJECTION","Try the provided DROP TABLE rejection test; show it is blocked.")]
for i,(n,h,d) in enumerate(demos):
    y=2.15+i*.92; shape(s,.9,y,11.45,.7,PANEL,line=RGBColor(38,55,82)); txt(s,1.15,y+.2,.55,.25,n,12,CYAN,True); txt(s,1.9,y+.18,2.4,.28,h,12,WHITE,True); txt(s,4.25,y+.18,7.6,.32,d,11,MUTED)
shape(s,.9,6.15,11.45,.48,RGBColor(12,35,43)); txt(s,1.15,6.27,10.9,.22,"PII SHIELD · INNOVEX  |  Thank you",12,GREEN,True,PP_ALIGN.CENTER)
footer(s,7)

prs.save(OUT)
print(OUT)
